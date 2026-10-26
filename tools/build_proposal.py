# -*- coding: utf-8 -*-
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1]
FONT = "맑은 고딕"
MONO = "Consolas"
NAVY = "1F3A5F"
BLUE = "2E75B6"
LIGHT = "EAF1F8"
LIGHT2 = "F5F8FC"
GRAY = "6B7785"
RED = "C0392B"
GREEN = "1E8449"
ORANGE = "CA6F1E"
CONTENT_W = 17.0  # cm

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)


# ---------- low-level helpers ----------
def rgb(h):
    return RGBColor.from_string(h)


def set_rfonts(rPr, font):
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(a)) is not None:
            del rFonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(a), font)
    rFonts.set(qn("w:eastAsia"), FONT)


def style_font(style, font=FONT, size=None, bold=None, color=None, italic=False):
    style.font.name = font
    set_rfonts(style.element.get_or_add_rPr(), font)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = rgb(color)
    style.font.italic = italic


P_SUCC_AFTER_SHD = ("w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                    "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
                    "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                    "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                    "w:pPrChange")


def shade_par(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.insert_element_before(shd, *P_SUCC_AFTER_SHD)


def border_par(p, sides, color, sz=12, space=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), str(sz))
            b.set(qn("w:space"), str(space))
            b.set(qn("w:color"), color)
            pBdr.append(b)
    pPr.insert_element_before(pBdr, "w:shd", *P_SUCC_AFTER_SHD)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection",
                               "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers",
                               "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")


def table_borders(table, color="BFC8D3", sz=4, inner=True, none=False):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        if none or (not inner and side.startswith("inside")):
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
        b.append(e)
    tblPr.insert_element_before(b, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                                "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    trPr.insert(0, c)


def add_rich(p, text, size=None, color=None, bold=False, font=FONT, italic=False):
    """Add text with **bold** markup to paragraph p."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        r = p.add_run(part)
        r.font.name = font
        set_rfonts(r._element.get_or_add_rPr(), font)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        r.bold = b
        r.italic = italic
    return p


def para(text="", size=None, color=None, bold=False, align=None, after=6, before=0,
         italic=False, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_rich(p, text, size=size, color=color, bold=bold, italic=italic)
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return p


def bullet(text, level=1, after=3):
    style = "List Bullet" if level == 1 else "List Bullet 2"
    return para(text, style=style, after=after)


def h1(num, title, new_page=False):
    p = doc.add_paragraph(style="Heading 1")
    add_rich(p, f"{num}  {title}", color="FFFFFF", bold=True, size=15)
    shade_par(p, NAVY)
    border_par(p, ("top", "left", "bottom", "right"), NAVY, sz=4, space=4)
    p.paragraph_format.page_break_before = new_page
    if not new_page:
        p.paragraph_format.space_before = Pt(22)
    return p


def h2(title):
    p = doc.add_paragraph(style="Heading 2")
    add_rich(p, title, color=NAVY, bold=True, size=12)
    border_par(p, ("bottom",), "C9D6E3", sz=6, space=2)
    return p


def h3(title, color=BLUE):
    p = doc.add_paragraph(style="Heading 3")
    add_rich(p, title, color=color, bold=True, size=10.5)
    return p


def callout(label, text, color=BLUE, fill=LIGHT):
    p = doc.add_paragraph()
    if label:
        add_rich(p, label + "  ", bold=True, color=color)
    add_rich(p, text)
    border_par(p, ("left",), color, sz=24, space=8)
    shade_par(p, fill)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.3)
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    return p


def code(lines, size=8.5):
    if doc.paragraphs and doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    for i, ln in enumerate(lines):
        p = doc.add_paragraph()
        add_rich(p, ln if ln else " ", size=size, font=MONO)
        shade_par(p, "F2F4F7")
        pf = p.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.3)
        pf.keep_with_next = i < len(lines) - 1
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def cell_write(cell, text, size=9, bold=False, color=None, align=None):
    cell.paragraphs[0].text = ""
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        add_rich(p, ln, size=size, bold=bold, color=color)
        pf = p.paragraph_format
        pf.space_after = Pt(1)
        pf.space_before = Pt(1)
        pf.line_spacing = 1.15
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(headers, rows, widths, size=9, first_col_fill=LIGHT, first_col_bold=True,
          center_cols=(), header_fill=NAVY, zebra=True):
    assert abs(sum(widths) - CONTENT_W) < 0.05, (sum(widths), headers)
    n = len(widths)
    if doc.paragraphs and doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=(1 if headers else 0) + len(rows), cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    r0 = 0
    if headers:
        hr = t.rows[0]
        repeat_header(hr)
        cant_split(hr)
        for i, h in enumerate(headers):
            c = hr.cells[i]
            c.width = Cm(widths[i])
            cell_write(c, h, size=size, bold=True, color="FFFFFF", align="center")
            shade_cell(c, header_fill)
        r0 = 1
    for ri, row in enumerate(rows):
        tr = t.rows[r0 + ri]
        cant_split(tr)
        for ci, val in enumerate(row):
            c = tr.cells[ci]
            c.width = Cm(widths[ci])
            is_first = ci == 0 and first_col_fill
            cell_write(c, val, size=size, bold=(is_first and first_col_bold),
                       color=(NAVY if is_first else None),
                       align=("center" if ci in center_cols else None))
            if is_first:
                shade_cell(c, first_col_fill)
            elif zebra and ri % 2 == 1:
                shade_cell(c, LIGHT2)
    keep_table(t)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    return t


def keep_table(t):
    rows = t.rows
    for ri, row in enumerate(rows):
        if ri == len(rows) - 1:
            break
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def note(text):
    return para(text, size=8.5, color=GRAY, after=8, italic=False)


def flow(boxes, box_w, arrow_w, fill=LIGHT, title_color=NAVY):
    widths = []
    for i in range(len(boxes)):
        widths.append(box_w)
        if i < len(boxes) - 1:
            widths.append(arrow_w)
    assert abs(sum(widths) - CONTENT_W) < 0.05, sum(widths)
    t = doc.add_table(rows=1, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    table_borders(t, none=True)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    row = t.rows[0]
    cant_split(row)
    for i, w in enumerate(widths):
        c = row.cells[i]
        c.width = Cm(w)
        if i % 2 == 0:
            title, sub = boxes[i // 2]
            cell_write(c, [f"**{title}**", sub], size=8.5, align="center")
            for p in c.paragraphs[:1]:
                for r in p.runs:
                    r.font.color.rgb = rgb(title_color)
            shade_cell(c, fill)
        else:
            cell_write(c, "▶", size=11, color=BLUE, align="center")
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)


# ---------- styles ----------
normal = doc.styles["Normal"]
style_font(normal, size=10)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
    st = doc.styles[name]
    style_font(st, size=size, bold=True, color=NAVY)
    st.paragraph_format.keep_with_next = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(14)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(14)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(6)
doc.styles["Heading 3"].paragraph_format.space_before = Pt(8)
doc.styles["Heading 3"].paragraph_format.space_after = Pt(3)
for name in ("List Bullet", "List Bullet 2"):
    style_font(doc.styles[name], size=10)
    doc.styles[name].paragraph_format.line_spacing = 1.25

# ---------- header / footer ----------
sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
add_rich(hp, "SKKU AI CAMPUS 바이오헬스 AI 1기  |  [P1] See-Through CXR 기획안  |  오형준",
         size=8, color=GRAY)
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
border_par(hp, ("bottom",), BLUE, sz=6, space=4)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_field(p, instr):
    r = p.add_run()
    r.font.size = Pt(8.5)
    r.font.color.rgb = rgb(GRAY)
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "1"
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for e in (f1, it, f2, t, f3):
        r._element.append(e)


add_rich(fp, "– ", size=8.5, color=GRAY)
add_field(fp, "PAGE")
add_rich(fp, " –", size=8.5, color=GRAY)

# =====================================================================
# COVER
# =====================================================================
para("", after=0)
para("2026 성균관대학교 AI CAMPUS 바이오헬스 AI 교육과정 1기  ·  [P1] 개인 프로젝트 기획안",
     size=10, color=BLUE, bold=True, after=10)
p = para("", after=0)
border_par(p, ("top",), NAVY, sz=18, space=1)
para("See-Through CXR", size=30, color=NAVY, bold=True, after=2, before=10)
para("근거를 보여주는 흉부 X선 폐렴 3분류 AI", size=18, color=NAVY, bold=True, after=6)
para("정상 · 세균성 · 바이러스성 폐렴 판별, 판단 근거의 시각화·정량 검증, 보호자 눈높이 설명",
     size=11.5, color=GRAY, after=12)
p = para("", after=14)
border_par(p, ("bottom",), NAVY, sz=18, space=1)

table(None, [
    ["프로젝트명", "See-Through CXR — 설명 가능한 흉부 X선 폐렴 3분류 서비스"],
    ["과정 / 과제", "[P1] CXR 흉부 질환 다중 분류 프로젝트 (개인 프로젝트)"],
    ["작성자", "오형준"],
    ["수행 기간", "2026.10.13(화) ~ 2026.10.23(금)"],
    ["데이터", "Chest X-Ray Images (Pneumonia), Kermany et al. 2018 — 5,856장 · 약 1.24 GB"],
    ["분류 대상", "0 정상 · 1 세균성 폐렴 · 2 바이러스성 폐렴"],
    ["모델", "ResNet18 전이학습 — 특징 고정 → 미세조정 → 불균형 보정 → 증강"],
    ["전처리 비교", "원본 · 히스토그램 평활화 · CLAHE · LPF(가우시안) · HPF(언샤프 마스킹) — 성능과 주목 영역(LFR) 비교"],
    ["설명 가능성(XAI)", "Grad-CAM · 폐 영역 집중도(LFR) · 확신도 기반 판단 보류 · 근거 문장"],
    ["보호자 설명", "규칙 기반 문장 틀 — 분류 결과를 보호자가 이해할 수 있는 쉬운 말로 설명"],
    ["AWS 인프라", "Amazon S3 · SageMaker Training Job · SageMaker Endpoint · EC2 (Flask / Streamlit)"],
    ["작업 환경", "iMac (코드 작성 · 소규모 시운전 · 화면 개발) + AWS 서울 리전 (정식 학습 · 평가 · 서비스)"],
    ["AI 원칙", "Explainable · Verifiable · Understandable · Human-in-the-Loop · Reproducible"],
], [4.2, 12.8], size=9)
note("※ 데이터 수량은 실제 파일 집계값이며, 성능 목표·인스턴스·일정 등 그 밖의 수치는 기획 단계의 계획값입니다.")
para("", after=6)
para("기획안 v0.4 (제출본)  |  2026. 10. 14  |  오형준", size=10, color=NAVY, bold=True,
     align="right")

# =====================================================================
# TOC
# =====================================================================
p = para("목  차", size=20, color=NAVY, bold=True, align="center", after=6)
p.paragraph_format.page_break_before = True
p = para("", after=18)
border_par(p, ("bottom",), BLUE, sz=8, space=1)
toc = [
    ("01", "프로젝트 개요 및 배경"),
    ("02", "목표 · 핵심 질문 · 범위"),
    ("03", "데이터 분석 및 전처리 계획"),
    ("04", "모델링 및 실험 설계"),
    ("05", "고전 영상처리 전처리 비교 — P 시리즈"),
    ("06", "설명 가능성(XAI) 설계 — 블랙박스 해소 전략"),
    ("07", "보호자 설명 계층 — 규칙 기반 템플릿"),
    ("08", "AWS 시스템 아키텍처"),
    ("09", "서비스 화면 및 API 설계"),
    ("10", "평가 계획 및 제출 산출물"),
    ("11", "추진 일정 (WBS)"),
    ("12", "리스크 관리 · 한계 · 기대 효과"),
]
for n, t in toc:
    p = doc.add_paragraph()
    add_rich(p, n + "   ", size=12, bold=True, color=BLUE)
    add_rich(p, t, size=12, color="222222")
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.8)
p = para("부록  참고문헌", size=11, color=GRAY, after=0)
p.paragraph_format.left_indent = Cm(0.8)

# =====================================================================
# 01
# =====================================================================
h1("01", "프로젝트 개요 및 배경", new_page=True)
h2("1.1 추진 배경 및 필요성")
para("폐렴은 전 세계 5세 미만 소아의 감염성 질환 사망 원인 가운데 가장 큰 비중을 차지하는 질환이며(WHO), "
     "흉부 X선(CXR, Chest X-Ray)은 폐렴 진단에서 가장 먼저, 가장 널리 쓰이는 영상 검사입니다. 그러나 판독 결과는 "
     "판독의의 숙련도와 업무량에 따라 편차가 생기고, 판독 인력이 부족한 환경에서는 판독 지연이 곧 치료 지연으로 이어집니다.")
para("특히 폐렴의 원인이 **세균성인지 바이러스성인지는 치료 방향과 직결**됩니다. 세균성 폐렴은 항생제 치료가 필요하지만, "
     "바이러스성 폐렴에 쓰는 항생제는 효과가 없고 항생제 내성 문제를 키웁니다. 영상만으로 원인을 확정할 수는 없지만, "
     "영상 패턴에서 원인을 추정한 보조 정보는 임상 판단의 우선순위를 정하는 데 도움이 될 수 있습니다.")
callout("기존 AI의 한계:",
        "딥러닝 CXR 분류 모델은 높은 정확도를 보고하지만, 판단 근거를 제시하지 않는 블랙박스(Black-box)로 작동합니다. "
        "실제로 폐렴 검출 모델이 병변이 아니라 병원별 촬영 장비의 특성이나 영상 위 글자·마커 같은 '지름길(shortcut)'을 "
        "학습해, 다른 병원 데이터에서 성능이 떨어진 사례가 보고되었습니다(Zech et al., 2018). "
        "정확도 숫자만으로는 모델이 '맞는 이유로 맞혔는지' 알 수 없습니다.", color=RED, fill="FBEDEB")
callout("본 프로젝트의 해답:",
        "정상·세균성·바이러스성 폐렴을 3분류하면서, 모든 예측에 대해 ① 근거 영역(Grad-CAM)을 보여주고, "
        "② 그 근거가 실제 폐 안에 있는지 수치(폐 영역 집중도, LFR)로 검증하며, ③ 확신이 낮거나 근거가 의심스러우면 "
        "'판독의 확인 필요'로 판단을 보류하며, ④ 이 결과를 정해진 문장 틀로 보호자가 이해할 수 있는 "
        "쉬운 말로 풀어 전달하는 설명 가능한 판독 보조 서비스를 AWS 위에 구축합니다. "
        "또한 고전 영상처리 전처리(대비 강조 · LPF · HPF)가 성능과 모델의 주목 영역을 어떻게 바꾸는지 비교합니다.",
        color=GREEN, fill="EAF5EE")

h2("1.2 핵심 가치")
bullet("**설명 가능성(Explainable):** 예측마다 히트맵·클래스별 확률·근거 문장을 함께 제시")
bullet("**이해 가능성(Understandable):** 같은 근거를 의료진에게는 수치로, 보호자에게는 쉬운 말로 — 받는 사람에 맞춘 설명")
bullet("**검증 가능성(Verifiable):** 설명을 '보여주는 것'에서 끝내지 않고, 그 설명이 타당한지를 폐 영역 집중도(LFR)로 정량 평가")
bullet("**안전성(Human-in-the-Loop):** 불확실한 사례는 자동 판정하지 않고 판독의에게 넘김 — AI는 보조, 최종 판단은 의료진")
bullet("**재현성(Reproducible):** 고정된 분할 CSV·난수 시드·S3 모델 주소로 누가 실행해도 같은 결과")

h2("1.3 기획 원칙 — 작게, 끝까지, 측정 가능하게")
para("본 프로젝트는 약 7일간의 개인 프로젝트입니다. 범위를 '3분류 모델 → 전처리 비교 → 설명 가능성 → 보호자 설명 → AWS 배포'라는 한 줄기로 한정하고, "
     "모든 목표에 측정 지표와 성공 기준을 붙였습니다. 또한 가이드 원칙에 따라 **실제 집계값과 계획값을 구분**하여 기록하고, "
     "결과 보고서에서는 계획값 옆에 실측값을 나란히 제시합니다. 데이터 분할·비교 실험·파일 규칙은 과정의 "
     "AWS 프로젝트 가이드를 따릅니다.")

# =====================================================================
# 02
# =====================================================================
h1("02", "목표 · 핵심 질문 · 범위")
h2("2.1 핵심 질문과 목표")
para("프로젝트 가이드가 제시한 네 가지 분석 질문을 그대로 프로젝트 목표로 연결하고, 서비스 배포와 보호자 설명 목표를 추가했습니다.")
table(["핵심 질문", "목표", "측정 방법", "성공 기준 (계획값)"], [
    ["Q1. 정상과 두 종류의 폐렴을 얼마나 구분하는가?", "3분류 모델 학습·비교",
     "Macro-F1, 클래스별 Recall, 혼동행렬", "테스트 Macro-F1 ≥ 0.75\n폐렴 민감도 ≥ 0.95"],
    ["Q2. 세균성과 바이러스성을 혼동하는 영상은 무엇인가?", "오분류 사례 유형화",
     "혼동행렬, 오분류 갤러리, 히트맵 비교", "세균↔바이러스 오분류 전수 검토, 유형 3개 이상 도출"],
    ["Q3. 모델이 폐 영역을 참고하는가?", "판단 근거의 타당성 정량 검증",
     "Grad-CAM + 폐 영역 집중도(LFR)", "정답 예측의 평균 LFR ≥ 0.70"],
    ["Q4. 다른 컴퓨터에서도 같은 결과를 재현할 수 있는가?", "재현 가능한 파이프라인",
     "같은 모델·입력으로 SageMaker Notebook 평가와 Endpoint 예측 비교", "예측 라벨 일치율 100%\n확률 차이 < 1e-4"],
    ["Q5. 학습한 모델을 화면과 서버에 연결할 수 있는가?", "Endpoint + EC2 서비스 배포",
     "시연, 응답 시간 측정", "영상 1장 응답 ≤ 3초\n(Grad-CAM 포함)"],
    ["Q6. 결과를 보호자가 이해할 수 있는 말로 정확하게 전달할 수 있는가?", "규칙 기반 보호자용 설명",
     "문장 틀 단위 테스트, 10건 이해도 평가", "문장 오류 0건\n이해도 평균 ≥ 4.0 / 5"],
    ["Q7. 고전 전처리는 성능과 모델이 보는 곳을 바꾸는가?", "전처리 비교 실험 (P 시리즈)",
     "영상 통계 · FFT, 검증 Macro-F1 · Recall · LFR", "전처리별 비교표 완성\n검증 기준으로 최종 전처리 선정"],
], [4.6, 3.6, 4.4, 4.4], size=8.5)
note("※ 성공 기준은 기획 단계의 목표치이며, 결과 보고서에서 실측값과 나란히 기록합니다. "
     "폐렴 민감도 = 세균성·바이러스성을 '폐렴'으로 묶었을 때의 Recall.")

h2("2.2 프로젝트 범위")
table(["구분", "포함 (In Scope)", "제외 (Out of Scope)"], [
    ["데이터", "Kermany CXR 5,856장 (JPEG)", "DICOM 원본, 성인·타 기관 데이터"],
    ["과제", "정상 / 세균성 / 바이러스성 3분류", "병변 분할, 다른 흉부 질환(14종 등) 분류"],
    ["전처리", "HE, CLAHE, LPF, HPF (기법마다 설정 1개)", "전처리 파라미터 최적화, 딥러닝 기반 영상 향상"],
    ["설명", "Grad-CAM, LFR, 확신도·판단 보류, 규칙 기반 근거 문장 · 보호자용 문장 틀",
     "LLM 기반 설명 생성 (향후 확장), 진단서 · 처방 생성"],
    ["인프라", "S3, SageMaker Training / Endpoint, EC2 (Flask · Streamlit)",
     "실서비스 수준 인증, 다중 사용자, 오토스케일링"],
    ["활용", "교육·연구용 판독 보조 프로토타입", "임상 진단 기기로서의 사용"],
], [2.6, 7.2, 7.2])

# =====================================================================
# 03
# =====================================================================
h1("03", "데이터 분석 및 전처리 계획")
h2("3.1 데이터 개요")
table(None, [
    ["데이터셋", "Chest X-Ray Images (Pneumonia) — Kermany et al., 2018 (Cell)"],
    ["출처", "중국 광저우 여성아동의료센터, 1~5세 소아의 흉부 X선 (전후면, AP)"],
    ["형식", "JPEG, 회색조, 영상마다 해상도 상이"],
    ["저장 위치", "s3://aws-say3-dataset1/pneumonia-xray/chest_xray/ (원천, 읽기 전용)\n"
                  "배포 압축 파일 chest_xray.zip 약 1.24 GB"],
    ["규모", "5,856장 — 로컬 파일과 S3 파일의 파일명·크기·체크섬 일치 (가이드 확인값)"],
], [3.4, 13.6])

h2("3.2 클래스 구성 (실제 집계값)")
table(["원래 분할", "정상", "세균성 폐렴", "바이러스성 폐렴", "합계"], [
    ["train (학습 후보)", "1,341", "2,530", "1,345", "5,216"],
    ["val (기존 검증)", "8", "8", "0", "16"],
    ["test (최종 테스트)", "234", "242", "148", "624"],
    ["전체", "1,583", "2,780", "1,493", "5,856"],
], [4.6, 3.1, 3.1, 3.1, 3.1], center_cols=(1, 2, 3, 4))
bullet("**검증셋 사용 불가:** 기존 val은 16장뿐이고 바이러스성이 0장 → 검증셋을 새로 구성해야 함")
bullet("**클래스 불균형:** train에서 세균성(2,530)이 정상·바이러스성의 약 1.9배")
bullet("**테스트셋 편차:** 바이러스성이 148장으로 가장 적어 바이러스성 지표의 변동 폭이 클 수 있음")

h2("3.3 목록표(manifest)와 라벨 생성 규칙")
para("모델이 어떤 파일을 읽었는지 추적하고 같은 평가를 반복하기 위해, 학습 전에 영상 5,856장의 목록표(manifest.csv)를 먼저 만듭니다.")
table(["열 이름", "기록할 값"], [
    ["relative_path", "데이터 루트 기준 영상 상대경로"],
    ["label_id / label_name", "0 · 1 · 2와 클래스 이름"],
    ["source_split / split", "원래 폴더(train · val · test)와 프로젝트에서 확정한 분할"],
    ["patient_group", "확인 가능한 동일 환자 그룹"],
    ["width / height / channels", "영상 크기와 원래 채널 수"],
    ["sha256 / readable", "파일 해시와 실제 읽기 성공 여부"],
], [4.6, 12.4], size=8.5)
para("폴더는 NORMAL / PNEUMONIA의 이진 구조이므로, 폐렴 폴더 안에서는 파일명의 키워드로 3개 라벨을 만듭니다.")
table(["라벨", "코드 이름", "화면 표시", "생성 규칙"], [
    ["0", "NORMAL", "정상", "NORMAL 폴더의 영상"],
    ["1", "BACTERIAL_PNEUMONIA", "세균성 폐렴", "PNEUMONIA 폴더 + 파일명에 'bacteria'"],
    ["2", "VIRAL_PNEUMONIA", "바이러스성 폐렴", "PNEUMONIA 폴더 + 파일명에 'virus'"],
], [1.6, 5.0, 3.4, 7.0], center_cols=(0,))
bullet("라벨을 정할 수 없는 파일은 오류로 처리하고 원인을 확인 — 완료 기준: 라벨 누락 0건, 읽기 실패 내역 확인, 클래스별 수가 3.2 표와 일치")
bullet("파일명은 **라벨 생성에만** 사용하고 모델 입력에는 절대 넣지 않음")

h2("3.4 데이터 분할 전략")
para("기존 train 5,216장 안에서 **환자 단위·클래스 층화**로 학습 약 80% / 검증 약 20%를 나눕니다. "
     "기존 val 16장은 바이러스성이 없어 주 검증셋으로 쓰지 않고 화면·입출력 점검용으로만 표시합니다. "
     "기존 test 624장은 설정을 고르기 전에 고정하며, 모델 선택이나 임계값 결정에는 사용하지 않습니다(최종 1회 평가).")
table(["분할", "용도", "규모", "비고"], [
    ["학습 (train)", "가중치 학습", "약 4,170장 (예상)", "기존 train에서 환자 단위 층화 분할"],
    ["검증 (val)", "에폭·실험 선택, 온도 보정, 보류 임계값 결정", "약 1,040장 (예상)", "기존 train에서 환자 단위 층화 분할"],
    ["테스트 (test)", "최종 성능 보고 (1회)", "624장 (실제)", "기존 test 그대로 고정"],
    ["점검용", "화면·입출력 동작 확인", "16장 (실제)", "기존 val — 성능 평가에 사용 안 함"],
], [3.0, 6.6, 3.4, 4.0])
note("※ 환자별 영상 수가 달라 최종 비율은 정확히 80 : 20이 아닐 수 있습니다.")
bullet("**환자 단위 분할:** 파일명의 환자 식별자(예: person1_bacteria_1.jpeg → person1)로 묶어 같은 환자의 영상이 "
       "학습과 검증에 함께 들어가지 않게 함 (StratifiedGroupKFold, seed 42). 정상(IM-xxxx)과 폐렴(personN)의 번호 체계가 "
       "같다고 가정하지 않고, 식별 규칙의 한계를 문서로 남김")
bullet("**중복 영상 점검:** 환자 그룹과 별도로 sha256 해시로 같은 파일이 여러 분할에 있는지 확인 (해시 일치와 동일 환자는 다른 조건)")
bullet("**완료 조건 검사:** 세 클래스가 학습·검증·테스트에 모두 있고, 환자 그룹과 동일 파일이 분할을 가로질러 겹치지 않는지 "
       "점수 계산 전에 코드로 검사")
bullet("**분할 고정:** split_manifest.csv(최종 분할)와 split_summary.csv(클래스별 영상 수·환자 수)를 S3에 저장해 모든 실험이 같은 분할을 사용")

h2("3.5 데이터 이슈와 대응")
table(["이슈", "내용", "대응"], [
    ["클래스 불균형", "세균성 비중이 가장 큼", "기준은 보정 없음 → 실험 C에서 train 빈도로만 계산한 클래스 가중치 적용, 주 지표는 Macro-F1"],
    ["해상도·명암 편차", "촬영 조건이 영상마다 다름", "기준(P0)은 224×224 리사이즈, 3채널 복제, ImageNet 정규화. 대비 강조 · 필터는 05장에서 비교.\n"
                                              "전처리 설정을 metadata.json에 저장해 Endpoint에도 똑같이 적용"],
    ["데이터 증강", "학습용 약 4.2천 장", "실험 D에서 한 가지 조건(예: ±10° 회전 또는 밝기 변화)만 학습셋에 추가.\n"
                                         "**좌우 반전은 사용하지 않음** (심장 등 해부학적 좌우 방향 보존)"],
    ["지름길 단서", "영상 위 글자·마커, 촬영 기기 차이", "EDA 샘플 갤러리로 육안 점검 → Grad-CAM·LFR로 사후 검증"],
    ["라벨 노이즈", "세균/바이러스 라벨의 확진 근거(배양 검사 등)가 공개되지 않음",
     "한계로 명시하고, 오분류 분석 시 라벨 의심 사례를 별도 표시"],
], [3.2, 5.4, 8.4])

h2("3.6 EDA 산출물")
bullet("분할·클래스별 분포 그래프, 환자당 영상 수 분포")
bullet("영상 크기·종횡비 분포, 클래스별 픽셀 밝기 히스토그램")
bullet("클래스별 샘플 갤러리(각 16장) — 글자·마커 등 지름길 단서 육안 점검용")

# =====================================================================
# 04
# =====================================================================
h1("04", "모델링 및 실험 설계")
h2("4.1 기준 모델 (Baseline)")
table(None, [
    ["백본", "ResNet18 — ImageNet 사전학습 가중치(resnet18_imagenet.pt)를 S3 weights 채널로 입력"],
    ["출력층", "fc 512 → 3 (정상 / 세균성 / 바이러스성)"],
    ["입력", "224 × 224, 3채널 (배열 모양 B × 3 × 224 × 224)"],
    ["손실 함수", "Cross-Entropy (Softmax 이전 점수 + 정수 정답). 클래스 가중치는 실험 C부터"],
    ["최적화", "Adam — 분류층 lr 0.001, 미세조정 lr 0.0001부터 비교 (가이드 시작 제안)"],
    ["배치 · 에폭", "배치 16, 최대 30 에폭, 검증 Macro-F1 최고 에폭 선택 (동률 시 규칙을 미리 고정)"],
    ["학습 환경", "SageMaker Training Job — CPU(ml.m5.xlarge)로 입출력 확인 후 허용 시 GPU(ml.g4dn.xlarge). "
                 "실행 제한 시간은 예비 실행 속도를 보고 확정"],
], [3.4, 13.6])

h2("4.2 비교 실험 설계")
table(["ID", "모델", "변경점", "확인할 것"], [
    ["A", "ResNet18", "특징 고정, 분류층만 학습 (배치 정규화 이동평균도 고정)", "사전학습 특징만으로 가능한 수준"],
    ["B", "ResNet18", "A + 마지막 잔차 블록(layer4)도 학습", "흉부 영상에 맞춘 특징 수정 효과"],
    ["C", "ResNet18", "B + 클래스 가중치", "적은 클래스(바이러스성)의 Recall 변화"],
    ["D", "ResNet18", "B + 증강 조건 1개", "촬영 변화에 대한 일반화 효과"],
    ["P 시리즈", "ResNet18", "B + 입력 전처리만 변경 (05장)", "고전 전처리가 성능 · LFR에 주는 영향"],
    ["E (선택)", "최종 모델", "폐 영역 마스킹 입력", "LFR이 낮을 때 지름길 제거 효과"],
], [2.0, 3.0, 5.0, 7.0], center_cols=(0,))

h2("4.3 공정 비교 원칙")
bullet("분할·라벨 순서·검증 선택 기준을 고정하고 **한 번에 하나의 조건만** 바꿈 (같은 split_manifest.csv, 전처리, 에폭 예산, seed 42)")
bullet("기준 실험(A~D)의 데이터와 계산을 먼저 확인한 뒤 P 시리즈와 선택 실험 E를 진행 — 다른 백본 모델 비교는 이번 범위에서 제외")
bullet("모델 선택·하이퍼파라미터 결정은 검증셋으로만 하고, 테스트셋은 최종 후보에 대해 한 번만 평가")
bullet("시간이 허락하면 최종 후보 2개를 시드 3개로 반복해 평균 ± 표준편차로 보고")
bullet("설정은 config.json에 보관하고 한 실험 안에서는 바꾸지 않음. 실험마다 별도 폴더에 model.tar.gz · history.csv · "
       "experiment_summary.json(설정·시드·장치·버전·선택 에폭·분할 파일 해시)을 같은 실험 ID로 저장")
bullet("작업 이름·입력 주소·설정 해시는 submission.json에 기록. 제한 시간에 멈춘(Stopped) 작업이나 실패(Failed) 작업을 완료로 기록하지 않음")

h2("4.4 최종 모델 선정 기준")
para("아래 순서로 우선순위를 둡니다. **성능이 비슷하면 근거가 더 타당한 모델을 고릅니다.**")
table(["순위", "기준", "이유"], [
    ["①", "검증 Macro-F1", "세 클래스를 고르게 맞히는지 (주 지표)"],
    ["②", "폐렴 민감도 (동률 시)", "폐렴을 정상으로 놓치는 오류가 가장 위험"],
    ["③", "정답 예측의 평균 LFR", "맞는 이유로 맞혔는지 (근거 타당성)"],
    ["④", "추론 지연·모델 크기", "Endpoint 서비스 비용과 응답 속도"],
], [1.6, 5.2, 10.2], center_cols=(0,))

# =====================================================================
# 05 preprocessing
# =====================================================================
h1("05", "고전 영상처리 전처리 비교 — P 시리즈")
callout("핵심 아이디어 — 레시피는 그대로, 재료 손질만 바꿉니다.",
        "모델과 학습 설정은 실험 B로 고정하고, 입력 영상의 전처리만 바꿔 성능과 모델이 주목하는 영역(LFR)이 어떻게 달라지는지 비교합니다. "
        "가이드의 '한 번에 하나의 조건만 바꾼다' 원칙을 그대로 따릅니다.", color=BLUE, fill=LIGHT)

h2("5.1 비교할 전처리")
table(["ID", "전처리", "하는 일", "흉부 X선에서 기대 · 우려", "고정 설정 (계획값)"], [
    ["P0", "원본 (기준)", "리사이즈 · 정규화만", "비교 기준 (실험 B 결과 재사용)", "—"],
    ["P1", "히스토그램 평활화 (HE)", "영상 전체의 밝기 분포를 고르게 펴서 대비를 키움", "전체 대비 증가, 뼈 · 테두리가 과하게 강조될 수 있음", "전역 평활화"],
    ["P2", "CLAHE", "작은 구역별로 대비를 키우되 과증폭을 제한", "폐 안의 옅은 음영이 잘 보일 것으로 기대", "clipLimit 2.0\n타일 8×8"],
    ["P3", "LPF\n(가우시안 블러)", "고주파 성분(노이즈 · 미세 무늬)을 줄임", "노이즈 감소, 바이러스성의 미세 간질 무늬도 약해질 우려", "σ = 1.5"],
    ["P4", "HPF\n(언샤프 마스킹)", "원본 + α × (원본 − 블러)로 경계 · 미세 무늬 강조", "미세 패턴 증가, 글자 · 테두리 · 노이즈도 강조 → LFR 저하 우려", "σ = 2, α = 1.0"],
    ["P5 (선택)", "3채널 조합", "[원본, CLAHE, HPF]를 RGB 3채널에 배치", "같은 영상을 3번 복제하는 대신 서로 다른 정보 3개를 입력", "P2 · P4 설정"],
], [1.8, 2.8, 4.2, 5.2, 3.0], size=8.5)
note("※ 각 기법은 설정값 하나만 사용하고 파라미터 탐색은 하지 않습니다. 설정은 config.json · metadata.json에 기록합니다.")

h2("5.2 진행 단계")
table(["단계", "내용", "장소", "완료 기준"], [
    ["1단계\n영상 비교", "같은 영상에 P0~P5를 적용한 갤러리, 클래스별 영상 통계(RMS 대비 · 엔트로피 · Sobel 에지 강도), "
                      "FFT 스펙트럼으로 저주파 · 고주파 에너지 변화 확인", "iMac\n(학습 없음)", "preprocess_gallery.png\npreprocess_stats.csv"],
    ["2단계\n모델 비교", "실험 B 설정에 전처리만 바꿔 P1~P4 (선택 P5) 학습, 검증 Macro-F1 · 클래스별 Recall · LFR 비교",
     "SageMaker\nTraining Job", "preprocess_comparison.csv"],
    ["3단계\n해석 (선택)", "ResNet18 첫 합성곱 층의 필터 시각화 — CNN이 스스로 학습한 저역 · 고역 통과형 필터 확인", "iMac", "first_layer_filters.png"],
], [2.4, 8.0, 2.8, 3.8], size=8.5)

h2("5.3 가설과 확인 방법")
table(["가설", "확인 지표"], [
    ["CLAHE(P2)는 폐 안 음영의 대비를 키워 바이러스성 Recall을 높인다", "클래스별 Recall, 혼동행렬"],
    ["LPF(P3)는 미세 무늬를 지워 세균성 ↔ 바이러스성 혼동을 늘린다", "세균 ↔ 바이러스 오분류 수"],
    ["HPF(P4)는 점수를 올리더라도 글자 · 테두리를 더 보게 만들어 LFR을 낮출 수 있다", "평균 LFR, Grad-CAM 비교"],
    ["전처리 효과는 작을 수 있다 — CNN의 첫 층이 비슷한 필터를 이미 학습하기 때문", "P0 대비 점수 차이, 첫 층 필터 그림"],
], [10.6, 6.4], size=8.5)

h2("5.4 원칙")
bullet("전처리는 증강이 아니라 **고정 변환** — 학습 · 검증 · 테스트 · Endpoint 모두에 똑같이 적용하고 설정을 metadata.json에 저장")
bullet("최종 전처리는 검증 Macro-F1로 고르고, 점수가 비슷하면 LFR이 높은 쪽을 선택. 테스트셋은 마지막에 한 번만 사용")
bullet("'P0 대비 차이 없음'도 유효한 결과로 보고하고, 3단계 필터 시각화로 그 이유를 설명")
bullet("일정이 부족하면 1단계(영상 비교)만으로도 결과 보고서에 포함")

# =====================================================================
# 05
# =====================================================================
h1("06", "설명 가능성(XAI) 설계 — 블랙박스 해소 전략")
callout("'블랙박스 없이'의 정의:",
        "모든 예측에 대해 사용자가 (1) 모델이 어디를 보았는지, (2) 그 근거가 타당한지, (3) 얼마나 확신하는지를 확인할 수 있고, "
        "(4) 개발자는 모델이 왜 틀렸는지 추적할 수 있는 상태.", color=NAVY, fill=LIGHT)

h2("6.1 4단계 근거 체계")
table(["단계", "질문", "방법", "사용자에게 보이는 것"], [
    ["① 시각적 근거", "어디를 보고 판단했나?", "Grad-CAM (예측 클래스 기준, 마지막 합성곱 층)",
     "원본 위 히트맵 오버레이, 투명도 조절"],
    ["② 근거 검증", "그 근거가 폐 안에 있나?", "폐 영역 집중도(LFR)\n= 폐 안 CAM 합 ÷ 전체 CAM 합",
     "LFR 수치와 '근거 신뢰 / 근거 의심' 표시"],
    ["③ 확신도", "얼마나 확신하나?", "Softmax 확률 + 온도 보정(Temperature Scaling)",
     "3클래스 확률 막대, 확정 / 보류 배지"],
    ["④ 오류 추적", "왜 틀렸나?", "혼동행렬 + 오분류 갤러리 + 히트맵 비교",
     "결과 보고서·발표 자료 (개발자용)"],
], [2.8, 3.4, 5.6, 5.2])

h2("6.2 Grad-CAM 적용")
bullet("라이브러리: pytorch-grad-cam")
bullet("대상 층: ResNet18의 마지막 합성곱 블록(layer4)")
bullet("예측 클래스(1순위)와 2순위 클래스의 히트맵을 모두 생성 → 세균성·바이러스성 근거 영역 비교")
bullet("**건전성 점검(Sanity Check):** 모델 가중치를 무작위로 바꿨을 때 히트맵이 크게 달라지는지 확인해, "
       "히트맵이 학습된 모델을 실제로 반영하는지 검증 (Adebayo et al., 2018)")

h2("6.3 폐 영역 집중도(LFR, Lung Focus Ratio) — 핵심 아이디어")
p = para("LFR  =  Σ ( CAM × 폐 마스크 )  ÷  Σ ( CAM )", size=12, bold=True, color=NAVY,
         align="center", before=4, after=2)
para("0 ~ 1 사이 값. 1에 가까울수록 폐 안의 영역을 근거로 판단했다는 뜻", size=9, color=GRAY,
     align="center", after=8)
bullet("**폐 마스크:** 공개 사전학습 폐 분할 모델(torchxrayvision의 해부학 구조 분할 모델)로 좌·우 폐 마스크 생성. "
       "소아 영상에서의 품질은 30장 육안 검수로 확인하고, 품질이 낮으면 폐 영역 바운딩 박스로 근사")
bullet("**해석:** LFR < 0.5이면 '폐 밖 근거 의심'으로 표시하고, 해당 사례를 지름길 학습 점검 대상으로 수집")
bullet("**활용:** ① 모델 간 비교 지표 ② 판단 보류 조건 ③ 정답·오답 사례 간 LFR 분포 비교")

h2("6.4 판단 보류 (Human-in-the-Loop)")
para("아래 조건 중 하나라도 해당하면 AI가 판정을 확정하지 않고 **'판독의 확인 필요'**로 표시합니다.")
table(["조건", "기준 (계획값)", "의미"], [
    ["최대 확률이 낮음", "보정 후 최대 확률 < τ (검증셋에서 결정, 예: 0.60)", "모델이 헷갈림"],
    ["1·2순위 차이가 작음", "p1 − p2 < 0.15", "세균성/바이러스성 경계 사례"],
    ["근거 의심", "LFR < 0.5", "폐 밖을 보고 판단했을 가능성"],
], [4.0, 7.4, 5.6])
para("결과 보고서에는 보류 비율(coverage)과 보류하지 않은 사례의 정확도를 함께 제시해, "
     "'얼마나 넘기면 얼마나 정확해지는가'를 곡선으로 보여줍니다.")

h2("6.5 근거 문장 생성 (규칙 기반)")
para("LLM을 쓰지 않고 수치와 위치 정보로 템플릿 문장을 만들어, 근거가 과장되거나 지어지는 일을 막습니다.")
code([
    "AI 판단 : 세균성 폐렴 (확률 0.82, 확정)",
    "근거 영역: 환자 기준 우측 폐 하부 (히트맵 활성의 78%가 폐 영역 안)",
    "2순위   : 바이러스성 폐렴 (0.13)",
], size=9)
bullet("위치는 CAM 무게중심을 폐 마스크 기준 좌/우 × 상/중/하로 변환해 표시")
bullet("영상의 왼쪽이 환자의 오른쪽인 방사선 영상 관례를 반영해 '환자 기준'으로 표기")
note("※ 위 예시의 수치는 화면 설계용 가상값입니다.")

h2("6.6 (탐색) 영상 패턴 가설 검증")
para("일반적으로 세균성 폐렴은 국소적인 경화(consolidation), 바이러스성 폐렴은 양측성·간질성 음영이 상대적으로 흔하다고 알려져 있습니다. "
     "이를 바탕으로 '세균성 예측의 히트맵은 한쪽 폐에 집중되고, 바이러스성 예측의 히트맵은 양쪽 폐에 넓게 퍼질 것'이라는 가설을 세우고, "
     "좌·우 폐의 CAM 비율과 활성 면적으로 검증합니다. 가설이 맞지 않더라도, 모델이 의학적 지식과 다른 근거를 쓰고 있음을 보여주는 결과로 활용합니다.")

# =====================================================================
# 06 guardian template
# =====================================================================
h1("07", "보호자 설명 계층 — 규칙 기반 템플릿")
callout("설계 원칙 — 정해진 문장 틀만 사용합니다.",
        "보호자용 설명은 미리 작성한 문장 틀에 분류 결과(라벨 · 확률 · 위치 · 보류 여부)를 끼워 넣어 만듭니다. "
        "틀에 없는 말은 나올 수 없으므로 사실과 다른 설명이나 과장된 표현이 생기지 않고, 같은 입력에는 항상 같은 문장이 나옵니다.",
        color=ORANGE, fill="FEF5E7")

h2("7.1 왜 보호자용 설명인가")
para("이 데이터는 1~5세 소아 영상이므로, 결과를 실제로 전달받는 사람은 환자 본인이 아니라 **보호자**입니다. "
     "'Grad-CAM', 'LFR 0.78', '세균성 폐렴 확률 0.82' 같은 표현은 의료진에게는 근거가 되지만 보호자에게는 또 다른 블랙박스입니다. "
     "같은 근거를 받는 사람에 맞는 언어로 바꿔 전달하는 것까지가 '블랙박스 없는' 서비스의 마지막 단계라고 보았습니다.")
table(["대상", "설명 방식", "생성 방법"], [
    ["의료진", "확률 · LFR · 근거 위치 등 수치 중심", "규칙 기반 근거 문장 (6.5절)"],
    ["보호자", "쉬운 말 4단 구성, 수치를 말로도 풀어 씀", "보호자용 문장 틀 (본 장)"],
], [3.4, 8.2, 5.4])

h2("7.2 문장 틀 설계")
code([
    "[확정된 경우]",
    "① AI가 사진을 분석한 결과, {질환명}일 가능성이 {가능성 표현}(약 {확률}%).",
    "② 그림에서 붉게 표시된 곳은 AI가 주로 살펴본 부분으로, 아이의 {위치}예요.",
    "③ 이 결과는 진단이 아니에요. {2순위 질환명}일 가능성(약 {2순위 확률}%)도 남아 있어요.",
    "④ 정확한 진단과 치료는 담당 의사 선생님이 결정해요. 궁금한 점은 의료진에게 물어보세요.",
    "",
    "[판단 보류인 경우]",
    "① AI가 이 사진만으로는 판단하기 어려워요. 의사 선생님의 확인이 필요해요.",
    "④ (위와 같은 안내 문장)",
], size=8.5)
note("※ ③의 2순위 문장은 2순위 확률이 기준 이상일 때만 표시합니다. 위 예시의 값은 설계용 가상값입니다.")

h2("7.3 값 변환 규칙 (계획값)")
table(["입력 값", "변환 규칙", "예시"], [
    ["최대 확률", "≥ 0.70 → '가능성이 높게 나왔어요' / 보류 기준 τ(약 0.60) ~ 0.70 → '가능성이 있어요' / τ 미만 → 보류 문장",
     "0.82 → 약 82%, 가능성이 높게 나왔어요"],
    ["2순위 확률", "≥ 0.10일 때만 ③의 2순위 문장 표시", "0.13 → 바이러스성 폐렴일 가능성(약 13%)"],
    ["히트맵 위치", "CAM 무게중심을 폐 마스크 기준 좌/우 × 상/중/하로 변환, 환자 기준으로 표기", "아이의 오른쪽 폐 아래쪽"],
    ["판단 보류", "판단 보류 조건(6.4절)에 해당하면 질환명 없이 보류 문장만 표시", "—"],
    ["정상 라벨", "'이상 소견이 뚜렷하지 않아요'로 표현하고, '질환이 없다'고 단정하지 않음", "—"],
], [2.8, 9.0, 5.2], size=8.5)

h2("7.4 안전 원칙")
bullet("문장 틀과 변환 규칙만 사용 — 입력에 없는 소견 · 질환 · 수치는 만들어질 수 없음")
bullet("④의 안내 문장(진단이 아니며 의료진이 결정한다는 내용)은 모든 경우에 고정으로 포함")
bullet("'확진', '틀림없이' 같은 단정 표현, 치료법 · 약 이름, 불안을 키우는 표현은 문장 틀에 넣지 않음")
bullet("문장 틀은 버전(template_v1)으로 관리하고, 응답 로그에 템플릿 버전을 함께 기록")

h2("7.5 품질 확인")
table(["지표", "방법", "목표 (계획값)"], [
    ["문장 생성 검사", "모든 라벨 × 확정 / 보류 × 2순위 표시 여부 조합에 대한 단위 테스트", "오류 0건"],
    ["사실 일치", "문장 속 숫자 · 질환명 · 위치를 입력값과 대조 (구조상 일치하며 테스트로 확인)", "100%"],
    ["이해도 · 친절성", "10건을 5점 척도로 평가 (본인 + 동료, 가능하면 비전공자 1명)", "평균 ≥ 4.0"],
], [3.6, 9.4, 4.0], size=8.5)

h2("7.6 향후 확장 — 생성형 AI 설명")
para("문장 틀 방식은 안전하지만 표현이 반복적입니다. 과정 후반의 생성형 AI 수업 이후에는, Amazon Bedrock의 LLM이 같은 입력 JSON만 받아 "
     "문장을 자연스럽게 다듬고, 문장 틀 결과와 숫자 · 질환명이 일치할 때만 표시하는 방식으로 확장할 수 있습니다. "
     "이번 P1 범위에는 포함하지 않습니다.")

# =====================================================================
# 07
# =====================================================================
h1("08", "AWS 시스템 아키텍처")
h2("8.1 전체 구성")
h3("학습 경로")
flow([("학생 PC", "코드 · 분할 CSV\n작업 제출"),
      ("Amazon S3", "원천 데이터\n분할 · 결과 저장"),
      ("SageMaker\nTraining Job", "train.py\n학습 · 검증 평가"),
      ("Amazon S3", "model.tar.gz\n학습 이력 CSV")], 3.6, 0.867)
h3("서비스 경로")
flow([("Amazon S3", "model.tar.gz"),
      ("SageMaker\nEndpoint", "예측 + Grad-CAM\n+ LFR"),
      ("EC2 · Flask", "/predict API\n로그 · 설명 문장"),
      ("EC2 · Streamlit", "결과 · 근거\n시각화 화면"),
      ("사용자", "브라우저 접속")], 2.8, 0.75, fill="EAF5EE", title_color=GREEN)

h2("8.2 작업 환경 역할 분담")
para("학생 PC는 작업 제출과 화면 접속에 쓰고, 정식 학습은 SageMaker가 만든 Linux 인스턴스에서 실행합니다. "
     "SageMaker 작업은 시작까지 몇 분이 걸리므로, 코드는 iMac에서 소규모로 먼저 시운전한 뒤 제출합니다.")
table(["작업", "iMac 24 (16GB)", "AWS (ap-northeast-2)"], [
    ["코드 작성 · 디버깅 · EDA 영상 확인", "주 작업 공간", "—"],
    ["train.py 시운전 (약 200장 × 1에폭, CPU)", "제출 전 오류 확인", "—"],
    ["manifest · 분할 생성", "코드 확인", "SageMaker Notebook에서 생성 후 S3 저장"],
    ["전처리 비교 1단계 (갤러리 · 통계 · FFT)", "주 작업 공간", "—"],
    ["정식 학습 (실험 A~D, P 시리즈, 선택 E)", "작업 제출 (AWS CLI · SageMaker SDK 2 계열)", "SageMaker Training Job"],
    ["최종 테스트 평가 · Grad-CAM · LFR", "코드 개발 (작은 샘플)", "SageMaker Notebook (evaluate.py)"],
    ["Streamlit 화면", "로컬 개발", "EC2에 최종 배포"],
    ["보호자 설명 문장 틀", "개발 · 단위 테스트", "EC2 (Flask 안에서 실행)"],
    ["추론 서비스", "—", "SageMaker Endpoint · EC2"],
], [5.4, 5.4, 6.2], size=8.5)
bullet("학습 코드에서 Mac의 MPS를 선택하지 않음 — 장치는 CUDA가 있으면 CUDA, 없으면 CPU로 두고 로그에 기록")
bullet("iMac에는 Python 3.11 / 3.12 프로젝트 전용 가상환경과 sagemaker>=2.250,<3, boto3를 두고, 실제 사용 버전을 requirements.txt에 고정")
bullet("AWS 접근 키는 aws configure로 로컬 자격 증명 파일에만 저장하고, 코드·문서·공유 폴더에는 남기지 않음")

h2("8.3 구성 요소별 역할")
table(["구성 요소", "AWS 서비스", "역할", "입력 → 산출물"], [
    ["데이터 저장", "Amazon S3", "원천 데이터·목록표·가중치·모델·평가 결과 보관", "원천 영상 → manifest, split, model.tar.gz, 평가 CSV"],
    ["모델 학습", "SageMaker Training Job", "train.py 실행, 실험별 Job 분리. 입력 채널 images · manifests · weights",
     "S3 입력 채널 → history.csv, model.tar.gz"],
    ["평가", "SageMaker Notebook", "evaluate.py: 선택 모델로 test 624장 1회 평가, Grad-CAM · LFR", "test_predictions.csv, metrics.csv"],
    ["모델 추론", "SageMaker Endpoint", "inference.py: 전처리 → 예측 → Grad-CAM → LFR", "영상 → JSON (라벨·확률·히트맵·LFR)"],
    ["API 서버", "EC2 + Flask", "Endpoint 호출, 입력 검증, 추론 로그 기록", "HTTP 요청 → JSON 응답"],
    ["보호자 설명", "EC2 + Flask (문장 틀)", "분류 결과 JSON → 의료진용 / 보호자용 문장", "JSON → 설명 문장"],
    ["화면", "EC2 + Streamlit", "업로드·결과·근거 시각화, 보호자용 설명 탭", "사용자 입력 → 결과 화면"],
    ["권한 · 로그", "IAM, CloudWatch", "최소 권한 역할, 학습·추론 로그 확인", "실행 로그"],
], [2.6, 3.4, 5.6, 5.4], size=8.5)

h2("8.4 S3 경로 설계")
code([
    "원천 (읽기 전용) : s3://aws-say3-dataset1/pneumonia-xray/chest_xray/",
    "출력 (학생 경로) : s3://aws-say3-student-dataset/students/<학생ID>/P1/cxr/",
    "  ├─ manifests/     # manifest.csv · split_manifest.csv · split_summary.csv",
    "  ├─ weights/       # resnet18_imagenet.pt · 폐 분할 모델 가중치 (미리 준비)",
    "  ├─ code/          # 학습 코드 묶음",
    "  ├─ experiments/   # 실험 ID별 model.tar.gz · history.csv · experiment_summary.json",
    "  ├─ reports/       # test_predictions.csv · metrics.csv · 오분류 · Grad-CAM 그림",
    "  ├─ preprocess/    # preprocess_gallery.png · preprocess_stats.csv · preprocess_comparison.csv",
    "  └─ templates/     # 보호자 설명 문장 틀 (template_v1)",
])
note("※ 학습 인스턴스는 인터넷에서 가중치를 직접 내려받지 않는 구조이므로, ImageNet 가중치와 LFR용 폐 분할 모델 가중치를 "
     "weights 경로에 미리 올려 둡니다. 원천 영상과 기존 분할 폴더는 덮어쓰지 않습니다.")

h2("8.5 인스턴스 계획 (계획값)")
table(["용도", "인스턴스", "선택 이유"], [
    ["학습", "ml.m5.xlarge (CPU)로 입출력 확인\n→ 허용 시 ml.g4dn.xlarge (GPU)", "가이드 권장 순서. 학습용·추론용 컨테이너는 공식 표에서 버전을 골라 고정"],
    ["추론", "ml.m5.large (CPU)\n→ 지연 시 ml.g4dn.xlarge", "ResNet18 단일 영상 추론은 CPU로 충분할 것으로 예상"],
    ["웹 서버", "EC2 t3.small", "Flask · Streamlit 경량 실행"],
], [2.6, 5.6, 8.8])
note("※ 실제 단가와 사용 시간은 실행 시점에 기록해 결과 보고서의 비용표에 반영합니다.")

h2("8.6 보안 · 비용 관리")
bullet("**최소 권한:** SageMaker 실행 역할은 원천 버킷 읽기 + 개인 버킷 읽기/쓰기만 허용")
bullet("**네트워크:** EC2 보안 그룹은 SSH(22)·Streamlit(8501) 포트를 본인 IP에서만 허용, Flask(5000)는 내부 호출 전용")
bullet("**데이터:** 업로드 영상은 추론 후 서버에 저장하지 않고, 로그에는 파일 해시와 결과만 기록")
bullet("**비용:** Endpoint는 테스트·시연 시간에만 가동, AWS Budgets 알림 설정")
bullet("**종료 체크리스트:** Endpoint → Endpoint Config → Model 삭제, EC2 중지·종료, 실행 중인 Training Job 없음 확인, "
       "S3 보관 범위 기록 — 단계별 콘솔 스크린샷과 시각을 종료 기록으로 제출")

# =====================================================================
# 07
# =====================================================================
h1("09", "서비스 화면 및 API 설계")
h2("9.1 Streamlit 화면 구성 (와이어프레임)")
t = doc.add_table(rows=5, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
table_borders(t, color="8FA6BF", sz=6)
wf = [5.2, 6.2, 5.6]
for i, w in enumerate(wf):
    t.columns[i].width = Cm(w)
for row in t.rows:
    cant_split(row)
    for i, c in enumerate(row.cells):
        c.width = Cm(wf[i])
top = t.rows[0].cells[0].merge(t.rows[0].cells[2])
cell_write(top, "**See-Through CXR** — 근거를 보여주는 폐렴 3분류   (교육·연구용, 진단 목적 아님)",
           size=9.5, color="FFFFFF", align="center")
shade_cell(top, NAVY)
heads = ["① 입력", "② 근거 시각화", "③ 판단 결과"]
bodies = [
    "X선 영상 업로드 (JPEG / PNG)\n샘플 영상 선택\n원본 미리보기",
    "Grad-CAM 히트맵 오버레이\n폐 마스크 윤곽선 표시\n투명도 슬라이더\n1·2순위 클래스 히트맵 전환",
    "판정 배지 (확정 / 판독의 확인 필요)\n3클래스 확률 막대\nLFR 게이지 (근거 신뢰도)\n근거 문장",
]
for i in range(3):
    cell_write(t.rows[1].cells[i], heads[i], size=9.5, bold=True, color=NAVY, align="center")
    shade_cell(t.rows[1].cells[i], LIGHT)
    cell_write(t.rows[2].cells[i], bodies[i], size=9, align="center")
gd = t.rows[3].cells[0].merge(t.rows[3].cells[2])
cell_write(gd, ["**④ 보호자용 설명 탭**",
                "쉬운 말 4단 설명 (요약 · 그림 보는 법 · 의미와 한계 · 다음 단계)  |  '의료진용 / 보호자용' 보기 전환  |  "
                "규칙 기반 문장 틀 (template_v1)"], size=9, align="center")
shade_cell(gd, "FEF5E7")
bot = t.rows[4].cells[0].merge(t.rows[4].cells[2])
cell_write(bot, "⑤ 추적 정보 — 모델 버전 · Training Job 이름 · S3 모델 주소 · 전처리 방식 · 문장 틀 버전 · 추론 시각 · 응답 시간 "
                "(어떤 모델이 낸 결과인지 추적 가능, Traceability)", size=8.5, color=GRAY, align="center")
shade_cell(bot, "F2F4F7")
keep_table(t)
doc.add_paragraph().paragraph_format.space_after = Pt(4)

h2("9.2 API 명세")
table(None, [
    ["경로", "POST /predict  (EC2 Flask → SageMaker Endpoint 호출)"],
    ["요청", "multipart/form-data — 흉부 X선 영상 1장"],
    ["응답", "JSON — 예측 라벨, 3클래스 확률, LFR, 보류 여부·사유, 근거 문장, 히트맵(PNG base64), 모델 버전, 응답 시간"],
    ["설명 API", "POST /explain — 입력: /predict 결과 JSON + 대상(guardian / clinician)\n"
                 "출력: 의료진용 또는 보호자용 문장, 문장 틀 버전"],
    ["오류", "400 (영상 형식 오류) · 503 (Endpoint 미가동) — 화면에 안내 문구 표시"],
], [2.6, 14.4])
code([
    "{",
    '  "label": "BACTERIAL_PNEUMONIA",  "label_ko": "세균성 폐렴",',
    '  "probs": {"NORMAL": 0.05, "BACTERIAL_PNEUMONIA": 0.82, "VIRAL_PNEUMONIA": 0.13},',
    '  "lfr": 0.78,  "abstain": false,  "abstain_reason": null,',
    '  "evidence_text": "환자 기준 우측 폐 하부 (히트맵 활성의 78%가 폐 영역 안)",',
    '  "cam_png": "<base64>",  "model_version": "C-resnet18-v1",  "latency_ms": 840',
    "}",
])
note("※ 응답 예시의 값은 설계용 가상값입니다.")

# =====================================================================
# 08
# =====================================================================
h1("10", "평가 계획 및 제출 산출물")
h2("10.1 평가 지표")
table(["지표", "정의", "목적", "목표 (계획값)"], [
    ["Macro-F1", "세 클래스 F1의 단순 평균", "주 지표 (불균형에 강건)", "≥ 0.75"],
    ["클래스별 Precision · Recall", "클래스마다 계산", "약한 클래스 파악", "바이러스성 Recall 별도 보고"],
    ["폐렴 민감도", "세균성+바이러스성을 '폐렴'으로 묶은 Recall", "폐렴을 놓치는 오류 방지", "≥ 0.95"],
    ["혼동행렬", "3 × 3 행렬", "세균↔바이러스 혼동 분석", "—"],
    ["AUROC (One-vs-Rest)", "클래스별 AUROC + 세 값의 평균(Macro AUROC)", "임계값과 무관한 판별력", "보고"],
    ["95% 신뢰구간", "환자 단위 부트스트랩 (같은 환자 영상을 함께 표집)", "점수의 불확실성", "주요 지표에 함께 보고"],
    ["ECE", "예측 확률과 실제 정답률의 차이", "확률의 신뢰도 (온도 보정 전후)", "보정 후 감소"],
    ["평균 LFR", "정답·오답 사례별 LFR 평균", "근거 타당성", "정답 사례 ≥ 0.70"],
    ["보류 비율 · 보류 외 정확도", "보류 사례 비율과 나머지 사례의 정확도", "HITL 효과 (trade-off)", "곡선 제시"],
    ["추론 지연", "요청부터 응답까지 시간", "서비스 품질", "≤ 3초"],
    ["전처리별 비교", "P0~P4의 검증 Macro-F1 · 클래스별 Recall · LFR", "고전 전처리의 효과", "비교표 제시"],
    ["보호자 문장 검사", "문장 틀 단위 테스트 (모든 라벨 · 보류 조합)", "설명 문장의 정확성", "오류 0건"],
    ["설명 이해도", "10건 5점 척도 평균", "보호자 눈높이 전달", "≥ 4.0"],
], [4.0, 5.0, 4.4, 3.6], size=8.5)
note("※ Accuracy는 단독 지표로 쓰지 않고 참고로만 보고합니다. 모든 점수는 영상 수·양성 수와 함께 제시하며, "
     "특정 클래스를 한 번도 예측하지 않아 분모가 0이 되는 경우의 처리 규칙을 미리 정하고 따로 기록합니다.")

h2("10.2 가이드 필수 제출물 대응표")
table(["가이드 필수 제출", "본 프로젝트 산출물", "형식"], [
    ["분할", "manifest.csv · split_manifest.csv · split_summary.csv (환자 그룹 · 해시 중복 0건 확인 포함)", "CSV"],
    ["학습 이력", "실험별 history.csv (epoch · train_loss · val_macro_f1 · val_macro_auroc · elapsed_seconds) + 곡선", "CSV, PNG"],
    ["평가 CSV", "test_predictions.csv (relative_path · group · true_label · pred_label · p_normal · p_bacterial · p_viral, "
               "+ LFR · 보류 여부) · metrics.csv", "CSV"],
    ["S3 모델 주소", "describe-training-job 응답의 ModelArtifacts 주소 + experiment_summary.json", "JSON"],
    ["AWS 구성", "아키텍처도, 인스턴스·IAM·보안 그룹 설정표", "문서"],
    ["시연", "Streamlit 시연 영상 (업로드 → 결과 → 근거 확인)", "MP4"],
    ["종료 기록", "종료 체크리스트 + 콘솔 스크린샷 + 시각", "문서"],
    ["(추가) 전처리 비교", "preprocess_gallery.png · preprocess_stats.csv · preprocess_comparison.csv", "PNG, CSV"],
    ["(추가) 보호자 설명", "문장 틀(template_v1), 단위 테스트 결과, 10건 이해도 평가표", "문서"],
], [3.6, 10.6, 2.8])

h2("10.3 재현성 확보")
bullet("seed 42 고정 (Python · NumPy · PyTorch) — 단, 같은 시드만으로 플랫폼 간 완전 동일 학습이 보장되지는 않음을 명시")
bullet("requirements.txt, SageMaker 컨테이너 버전, CUDA 인식 여부를 작업 로그와 experiment_summary.json에 기록")
bullet("config.json(학습 설정)과 metadata.json(전처리 설정)을 모델과 함께 저장해 Endpoint에서도 같은 전처리 사용")
bullet("**검증:** 같은 모델과 테스트 영상으로 SageMaker Notebook 평가 확률과 Endpoint 예측 확률을 비교해 일치율 보고 (Q4)")

# =====================================================================
# 09
# =====================================================================
h1("11", "추진 일정 (WBS)")
para("P1 기간은 2026.10.13(화) ~ 10.23(금)입니다. 10/15(목)는 '의료 AI 보안 규제' 수업, 10/16(금) 오전은 복습 시간이므로 "
     "실제 프로젝트 작업일은 약 7일입니다. 전처리 비교 1단계와 보호자 설명 문장 틀은 AWS 없이 iMac에서 학습 작업과 병행합니다.")
table(["날짜", "단계", "주요 작업", "산출물"], [
    ["10/13 (화)", "기획", "기획안 초안 작성, 과정 가이드 검토", "기획안 초안"],
    ["10/14 (수)", "기획 확정 · 데이터", "기획안 v0.4 제출, iMac 환경·AWS 접속 확인, manifest 생성, 환자 단위 분할, EDA",
     "기획안 v0.4\nmanifest · split"],
    ["10/15 (목)", "(보안 수업)", "수업 내용을 IAM 최소 권한·보안 그룹 설계에 반영", "보안 설정표"],
    ["10/16 (금)", "학습 착수", "가중치 S3 업로드, iMac 시운전, 실험 A 제출 / (iMac) 전처리 1단계 갤러리 · 통계 · FFT",
     "실험 A 결과\n전처리 갤러리"],
    ["10/19 (월)", "실험", "실험 B~D 실행, 검증 결과 비교 → P 시리즈(P1~P4) 제출 / (iMac) 보호자 문장 틀 작성",
     "실험별 결과 폴더\n문장 틀 v1"],
    ["10/20 (화)", "XAI · 평가", "P 시리즈 비교표, 최종 모델 · 전처리 선정(검증 기준), 테스트 1회 평가, 부트스트랩 신뢰구간, Grad-CAM·폐 마스크·LFR, 보류 임계값",
     "test_predictions\nmetrics · XAI 그림"],
    ["10/21 (수)", "배포", "inference.py 작성, Endpoint 배포와 예측 확률 검산, EC2 Flask · Streamlit 연결, 보호자 문장 틀 연동",
     "동작하는 서비스"],
    ["10/22 (목)", "검증 · 정리", "통합 테스트, 재현성 검증, 보호자 설명 이해도 평가(10건), 시연 녹화, 보고서·발표 자료 작성",
     "시연 영상\n보고서"],
    ["10/23 (금)", "발표 · 종료", "최종 발표, AWS 리소스 종료 및 종료 기록", "종료 기록"],
], [2.4, 2.6, 8.6, 3.4], size=8.5, center_cols=(0,))

h2("11.2 우선순위 (일정 지연 시 축소 순서)")
table(["우선순위", "항목"], [
    ["필수 (Must)", "manifest · 분할, 실험 A~C, 평가 CSV, Grad-CAM, Endpoint 배포, Streamlit 화면, 종료 기록"],
    ["중요 (Should)", "실험 D, P 시리즈 1 · 2단계, LFR, 판단 보류, 재현성 검증(Endpoint 검산), 부트스트랩 신뢰구간, 보호자 문장 틀"],
    ["선택 (Could)", "P5 3채널 조합, 첫 층 필터 시각화(3단계), 실험 E(폐 마스킹), 온도 보정, 다중 시드, 영상 패턴 가설, LLM 설명(향후 확장)"],
], [3.4, 13.6])

# =====================================================================
# 10
# =====================================================================
h1("12", "리스크 관리 · 한계 · 기대 효과")
h2("12.1 리스크와 대응")
table(["리스크", "영향", "대응"], [
    ["세균/바이러스 구분 성능이 낮음", "높음", "폐렴 민감도와 Macro-F1을 나눠 목표를 설정하고, 혼동 분석 자체를 핵심 결과로 제시"],
    ["히트맵이 폐 밖을 가리킴 (지름길 학습)", "중간", "발견 자체를 결과로 보고하고, 선택 실험 E(폐 마스킹 입력)로 개선 시도"],
    ["폐 분할 모델이 소아 영상에서 부정확", "중간", "30장 육안 검수, 실패 시 바운딩 박스 근사로 LFR 계산"],
    ["SageMaker 권한·쿼터 문제", "높음", "첫날 학생 경로 읽기·쓰기와 실행 역할 점검. GPU가 허용되지 않으면 CPU로 실험 A·B를 우선 진행하고 강사와 협의"],
    ["학습 중 가중치를 내려받지 못함", "중간", "ImageNet · 폐 분할 가중치를 S3 weights 경로에 미리 업로드"],
    ["Endpoint 추론 지연 (Grad-CAM 포함)", "낮음", "GPU 인스턴스로 전환하거나 히트맵 생성을 선택 옵션으로 분리"],
    ["전처리 효과가 없거나 미미함", "중간", "유효한 결과로 보고하고, 첫 층 필터 시각화로 'CNN이 필터를 스스로 학습한다'는 해석 제시"],
    ["HPF가 글자 · 노이즈를 강조해 지름길 학습 유발", "중간", "LFR과 Grad-CAM으로 확인하고, 최종 전처리 선택에 LFR을 함께 반영"],
    ["AWS 비용 초과", "중간", "시연 시간에만 Endpoint 가동, Budgets 알림, 종료 체크리스트"],
    ["일정 지연", "중간", "11.2 우선순위에 따라 선택 항목부터 축소"],
], [5.0, 1.6, 10.4], size=8.5, center_cols=(1,))

h2("12.2 한계 및 윤리적 고려")
bullet("**데이터 일반화:** 단일 기관·1~5세 소아 영상 → 성인이나 다른 병원 영상에 그대로 적용할 수 없음")
bullet("**라벨 신뢰도:** 세균성/바이러스성 라벨의 확진 근거가 공개되지 않아 라벨 오류 가능성이 있음")
bullet("**예측 대상의 의미:** 영상으로 데이터셋 라벨을 예측할 뿐, 세균·바이러스의 실제 원인을 확정하는 검사가 아님. "
       "NORMAL은 데이터셋의 정상 라벨이며 모든 질환이 없다는 판정으로 확대하지 않음")
bullet("**설명의 한계:** Grad-CAM은 저해상도 특징맵(ResNet18 기준 7×7) 기반이며, 병변의 정확한 경계나 임상 원인을 입증하지 않음 → "
       "'모델이 참고한 영역'으로만 해석하고 모델 버전과 대상 클래스를 함께 표기")
bullet("**쉬운 설명의 위험:** 쉬운 말로 바꾸는 과정에서 불확실성이 작게 전달되거나 보호자의 불안을 키울 수 있음 → "
       "확률·한계 문장과 의료진 상담 안내를 항상 포함하고, 불안을 키우는 표현은 금지")
bullet("**사용 목적:** 교육·연구용 프로토타입이며 진단 기기가 아님. 최종 판단은 의료진")
bullet("**개인정보:** 공개된 비식별 데이터셋만 사용, 업로드 영상은 저장하지 않음")

h2("12.3 기대 효과")
table(["구분", "기대 효과"], [
    ["기술", "설명을 '보여주는 것'에서 '검증하는 것'으로 — LFR로 XAI 결과의 타당성을 수치로 비교하는 틀 확보. "
            "고전 전처리와 CNN이 학습한 필터의 관계를 실험으로 확인"],
    ["보호자 소통", "AI 결과를 보호자 눈높이로 풀어 전달해, 의료진의 설명 부담을 덜고 보호자의 이해를 돕는 소통 보조 (가정)"],
    ["임상적 시사점", "판독 보조 시 근거 영역과 확신도를 함께 제시하고 불확실 사례를 선별해, 판독의가 확인할 우선순위를 정하는 데 도움 (가정)"],
    ["역량", "S3 → SageMaker → Endpoint → EC2 전 주기 MLOps 경험, 고전 영상처리 파이프라인 설계 경험. 이후 M7-1 '의료 AI 모델 검증 및 해석', "
            "P2 '뇌 MRI 세그멘테이션'(분할 마스크 활용)으로 확장"],
], [3.4, 13.6])

p = para("", after=4, before=10)
border_par(p, ("top",), BLUE, sz=6, space=1)
para("본 기획안은 v0.4 제출본입니다. 데이터 수량을 제외한 모든 수치는 계획값이며, 결과 보고서에서 실측값으로 갱신합니다. "
     "See-Through CXR은 AI가 판독의를 대체하는 것이 아니라, 판단의 근거와 불확실성을 투명하게 드러내 "
     "판독의의 확인을 돕는 보조 도구를 지향합니다.", size=9, color=GRAY)
para("— Explainable · Verifiable · Understandable · Human-in-the-Loop · Reproducible  |  오형준  |  2026.10.14",
     size=9, color=BLUE, bold=True, align="right")

# =====================================================================
# 참고문헌
# =====================================================================
h2("부록. 참고문헌")
refs = [
    "Kermany, D. S. et al. (2018). Identifying Medical Diagnoses and Treatable Diseases by Image-Based Deep Learning. Cell, 172(5), 1122–1131.",
    "Zech, J. R. et al. (2018). Variable generalization performance of a deep learning model to detect pneumonia in chest radiographs: A cross-sectional study. PLOS Medicine, 15(11).",
    "Selvaraju, R. R. et al. (2017). Grad-CAM: Visual Explanations from Deep Networks via Gradient-based Localization. ICCV.",
    "Adebayo, J. et al. (2018). Sanity Checks for Saliency Maps. NeurIPS.",
    "He, K. et al. (2016). Deep Residual Learning for Image Recognition. CVPR.",
    "Rajpurkar, P. et al. (2017). CheXNet: Radiologist-Level Pneumonia Detection on Chest X-Rays with Deep Learning. arXiv:1711.05225.",
    "Guo, C. et al. (2017). On Calibration of Modern Neural Networks. ICML.",
    "Cohen, J. P. et al. (2022). TorchXRayVision: A library of chest X-ray datasets and models. MIDL.",
    "World Health Organization. Pneumonia in children — Fact sheet.",
    "Zuiderveld, K. (1994). Contrast Limited Adaptive Histogram Equalization. Graphics Gems IV, Academic Press.",
    "Gonzalez, R. C., & Woods, R. E. (2018). Digital Image Processing (4th ed.). Pearson.",
]
for i, r in enumerate(refs, 1):
    p = para(f"[{i}] {r}", size=8.5, color="333333", after=3)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)

doc.save(OUT)
print("saved", OUT)
